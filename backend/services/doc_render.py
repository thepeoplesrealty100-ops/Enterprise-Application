"""
services/doc_render.py — shared DOCX/PDF rendering for generated documents.

Both routers/report_builder.py (CPENT-style pentest reports) and
routers/rfp.py (RFP responses) build a plain, format-agnostic list of
sections --

    [{"heading": str, "body": str | None, "table": {"columns": [...], "rows": [[...]]} | None}, ...]

-- and hand it to `to_docx()` / `to_pdf()` here, so page layout (fonts,
table styling, margins) lives in one place instead of being reimplemented
per-router. DOCX output is real and editable (python-docx) so an operator
can open a generated report in Word and strike or adjust a section rather
than being stuck with a flattened PDF; PDF (reportlab, pure-Python, no
native deps) is the polished client-delivery format.
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional


def to_docx(title: str, subtitle: Optional[str], sections: List[Dict[str, Any]]) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        if p.runs:
            p.runs[0].italic = True

    for sec in sections:
        doc.add_heading(sec["heading"], level=1)
        body = sec.get("body")
        if body:
            for para in str(body).split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
        table = sec.get("table")
        if table and table.get("rows"):
            cols = table["columns"]
            t = doc.add_table(rows=1, cols=len(cols))
            try:
                t.style = "Light Grid Accent 1"
            except KeyError:
                pass
            hdr = t.rows[0].cells
            for i, c in enumerate(cols):
                hdr[i].text = str(c)
            for row in table["rows"]:
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = "" if val is None else str(val)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(title: str, subtitle: Optional[str], sections: List[Dict[str, Any]]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    flow: List[Any] = [Paragraph(title, styles["Title"])]
    if subtitle:
        flow.append(Paragraph(subtitle, styles["Italic"]))
    flow.append(Spacer(1, 16))

    for sec in sections:
        flow.append(Paragraph(sec["heading"], styles["Heading2"]))
        body = sec.get("body")
        if body:
            for para in str(body).split("\n\n"):
                if para.strip():
                    flow.append(Paragraph(para.strip().replace("\n", "<br/>"), styles["BodyText"]))
                    flow.append(Spacer(1, 6))
        table = sec.get("table")
        if table and table.get("rows"):
            data = [table["columns"]] + [[("" if v is None else str(v)) for v in row] for row in table["rows"]]
            t = Table(data, hAlign="LEFT")
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            flow.append(t)
        flow.append(Spacer(1, 12))

    doc.build(flow)
    return buf.getvalue()
